"""Exporta a aba "Base Governo" da planilha BaseLicitacao.xlsx (valores calculados) para um JSON.

Somente leitura: trabalha numa CÓPIA da planilha (nunca abre o original, para não travar quem edita) e apaga
a cópia no fim. Usado pela carga automática (scripts/carga-planilha.cmd) -> scripts/sync-db.mjs.

Uso:
  python scripts/exportar_planilha.py --saida PASTA [--origem CAMINHO.xlsx]
  (origem padrão: variável PLANILHA_ORIGEM ou o caminho do OneDrive sincronizado do time)
Requer: pip install openpyxl
"""
import argparse
import datetime as dt
import json
import os
import shutil
import sys
import tempfile

import openpyxl

ORIGEM_PADRAO = r"C:\Users\CTC\Connectcom Teleinformatica\GOVERNO - Documentos\General\LICITAÇÕES\Controles\Oportunidades\BaseLicitacao.xlsx"
ABA = "Base Governo"


def celula(v):
    if isinstance(v, (dt.datetime, dt.date)):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, dt.time):
        return None
    if isinstance(v, str):
        v = v.strip()
        return v or None
    return v


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--saida", required=True, help="pasta onde gravar aba.json")
    ap.add_argument("--origem", default=os.environ.get("PLANILHA_ORIGEM", ORIGEM_PADRAO))
    a = ap.parse_args()

    if not os.path.isfile(a.origem):
        print(f"ERRO: planilha não encontrada: {a.origem} (OneDrive sincronizado? 'Sempre manter neste dispositivo'?)", file=sys.stderr)
        return 2
    st = os.stat(a.origem)
    print(f"origem modificada em {dt.datetime.fromtimestamp(st.st_mtime).isoformat(timespec='seconds')} | {st.st_size} bytes")
    if st.st_size < 10_000:
        print("ERRO: arquivo pequeno demais; pode ser um 'espaço reservado' do OneDrive (ainda não baixado).", file=sys.stderr)
        return 2

    os.makedirs(a.saida, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix="radar-copia-") as tmp:
        copia = os.path.join(tmp, "BaseLicitacao.xlsx")
        shutil.copy2(a.origem, copia)
        wb = openpyxl.load_workbook(copia, read_only=True, data_only=True)
        if ABA not in wb.sheetnames:
            print(f"ERRO: aba '{ABA}' não encontrada (abas: {wb.sheetnames})", file=sys.stderr)
            return 3
        linhas = [[celula(c) for c in r] for r in wb[ABA].iter_rows(values_only=True)]
        wb.close()

    # remove colunas totalmente vazias à direita e linhas totalmente vazias
    largura = max((max((i for i, c in enumerate(r) if c is not None), default=-1) for r in linhas), default=-1) + 1
    linhas = [r[:largura] for r in linhas]
    linhas = [r for r in linhas if any(c is not None for c in r)]
    cab = [str(c or "") for c in (linhas[0] if linhas else [])]
    if not linhas or not any(c.lower().startswith("status") for c in cab) or not any("pregão" in c.lower() or "pregao" in c.lower() for c in cab):
        print("ERRO: cabeçalho inesperado (faltam 'Status' e 'Data Pregão').", file=sys.stderr)
        return 4
    destino = os.path.join(a.saida, "aba.json")
    with open(destino, "w", encoding="utf-8") as f:
        json.dump(linhas, f, ensure_ascii=False)
    print(f"ok: {len(linhas) - 1} linhas de dados, {largura} colunas -> {destino}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
